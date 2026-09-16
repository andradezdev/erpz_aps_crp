# Copyright (c) 2026, ERPZ and contributors
# For license information, please see license.txt

import frappe
from frappe import _
from frappe.utils import now_datetime, getdate, get_datetime, flt, cint
from erpz_aps.engine.aps_engine import APSEngine
from erpz_aps.engine.chain_propagation import propagate_operation_change
from erpz_aps.engine.execution import execute_aps_schedule

@frappe.whitelist()
def run_aps_calculation(ticket_name):
    """Triggers the finite-capacity APS/CRP engine for a given ticket."""
    if not frappe.has_permission("APS Ticket", "write"):
        frappe.throw(_("Sem permissão para executar programação APS."))
        
    engine = APSEngine(ticket_name)
    result = engine.run()
    return result

@frappe.whitelist()
def recalculate_chain(ticket_name, op_name, new_start_datetime, new_workstation=None):
    """
    Called when an operation is dragged and dropped in the interactive Gantt chart.
    Propagates the movement forward across all dependent operations and Work Orders in chain.
    """
    if not frappe.has_permission("APS Ticket", "write"):
        frappe.throw(_("Sem permissão para ajustar programação."))
        
    res = propagate_operation_change(ticket_name, op_name, new_start_datetime, new_workstation)
    return res

@frappe.whitelist()
def approve_ticket(ticket_name):
    """Approves an APS ticket scenario for execution."""
    if not frappe.has_permission("APS Ticket", "write"):
        frappe.throw(_("Sem permissão para aprovar Ticket APS."))
        
    ticket = frappe.get_doc("APS Ticket", ticket_name)
    if ticket.status in ("Efetivado", "Cancelado"):
        frappe.throw(_("Ticket já finalizado não pode ser aprovado."))
        
    ticket.db_set({
        "status": "Aprovado",
        "approved_on": now_datetime(),
        "approved_by": frappe.session.user
    })
    frappe.db.commit()
    return {"status": "success", "message": _("Ticket APS aprovado com sucesso.")}

@frappe.whitelist()
def execute_ticket(ticket_name):
    """Synchronizes scheduled start/end dates directly to Work Orders and Job Cards."""
    if not frappe.has_permission("APS Ticket", "write"):
        frappe.throw(_("Sem permissão para efetivar Ticket APS."))
        
    res = execute_aps_schedule(ticket_name)
    return res

@frappe.whitelist()
def get_gantt_data(ticket_name, group_by="workstation"):
    """
    Returns structured data for the interactive Gantt chart:
    - tasks: list of operation bars
    - links: dependency arrows between predecessor and successor operations
    - resources: workstations
    """
    ops = frappe.get_all(
        "APS Scheduled Operation",
        filters={"aps_ticket": ticket_name},
        fields=[
            "name", "work_order", "mrp_ticket", "production_item", "item_name",
            "operation", "sequence_id", "workstation", "planned_start_time",
            "planned_end_time", "duration_mins", "setup_mins", "runtime_mins",
            "qty_to_produce", "predecessor_operation", "is_fixed", "is_adjusted",
            "has_conflict", "status", "delay_hours"
        ],
        order_by="planned_start_time ASC"
    )

    workstations = frappe.get_all("Workstation", fields=["name", "workstation_name"])
    ws_map = {w.name: w.workstation_name or w.name for w in workstations}

    tasks = []
    links = []

    # Map for building dependency arrows
    wo_ops_map = {}

    for o in ops:
        start_str = str(o.planned_start_time)
        end_str = str(o.planned_end_time)
        
        # Color coding
        progress = 1.0 if o.status == "Concluída" else (0.5 if o.status == "Em Execução" else 0.0)
        
        # Group label
        ws_label = ws_map.get(o.workstation, o.workstation)
        
        task_id = o.name
        task_label = f"[{o.work_order}] {o.operation} ({int(o.duration_mins)}m)"
        
        tasks.append({
            "id": task_id,
            "text": task_label,
            "start_date": start_str,
            "end_date": end_str,
            "duration": o.duration_mins,
            "progress": progress,
            "workstation": o.workstation,
            "workstation_label": ws_label,
            "work_order": o.work_order,
            "production_item": o.production_item,
            "item_name": o.item_name,
            "operation": o.operation,
            "sequence_id": o.sequence_id,
            "qty": o.qty_to_produce,
            "is_adjusted": o.is_adjusted,
            "has_conflict": o.has_conflict,
            "status": o.status,
            "delay_hours": o.delay_hours
        })

        # Track operations by Work Order and sequence
        key = (o.work_order, o.sequence_id)
        wo_ops_map[key] = task_id

    # Build dependency links (predecessor -> successor in the same Work Order)
    for o in ops:
        current_id = o.name
        current_seq = o.sequence_id
        # Find predecessor operation with highest sequence < current_seq in the same WO
        preds = [
            wo_ops_map[(o.work_order, s)] 
            for (w, s) in wo_ops_map.keys() 
            if w == o.work_order and s < current_seq
        ]
        if preds:
            pred_id = preds[-1]
            links.append({
                "id": f"{pred_id}_{current_id}",
                "source": pred_id,
                "target": current_id,
                "type": "0" # finish-to-start
            })

    return {
        "tasks": tasks,
        "links": links,
        "workstations": workstations
    }

@frappe.whitelist()
def get_crp_data(ticket_name, workstation=None):
    """Returns CRP Capacity vs Load data and chart series."""
    filters = {"aps_ticket": ticket_name}
    if workstation:
        filters["workstation"] = workstation

    rows = frappe.get_all(
        "APS Capacity Load",
        filters=filters,
        fields=[
            "workstation", "workstation_name", "period_date",
            "available_capacity_hours", "planned_load_hours",
            "capacity_balance_hours", "utilization_pct", "is_bottleneck",
            "overload_hours", "idle_hours", "operations_count"
        ],
        order_by="period_date ASC, workstation ASC"
    )

    dates = sorted(list(set(str(r.period_date) for r in rows)))
    
    # Aggregated daily capacity vs load
    daily_capacity = {d: 0.0 for d in dates}
    daily_load = {d: 0.0 for d in dates}
    
    for r in rows:
        d_str = str(r.period_date)
        daily_capacity[d_str] += flt(r.available_capacity_hours)
        daily_load[d_str] += flt(r.planned_load_hours)

    return {
        "dates": dates,
        "capacity_series": [round(daily_capacity[d], 1) for d in dates],
        "load_series": [round(daily_load[d], 1) for d in dates],
        "raw_rows": rows
    }

@frappe.whitelist()
def get_ticket_dashboard_metrics(ticket_name):
    """Returns high-level KPI cards data for the APS workbench."""
    ticket = frappe.get_doc("APS Ticket", ticket_name)
    
    total_ops = frappe.db.count("APS Scheduled Operation", {"aps_ticket": ticket_name})
    total_bottlenecks = frappe.db.count("APS Capacity Load", {"aps_ticket": ticket_name, "is_bottleneck": 1})
    delayed_count = frappe.db.count("APS Scheduled Operation", {"aps_ticket": ticket_name, "status": "Atrasada"})
    adjusted_count = frappe.db.count("APS Scheduled Operation", {"aps_ticket": ticket_name, "is_adjusted": 1})
    
    return {
        "status": ticket.status,
        "total_orders": ticket.total_orders or 0,
        "total_operations": total_ops,
        "total_scheduled_hours": ticket.total_scheduled_hours or 0.0,
        "average_utilization_pct": ticket.average_utilization_pct or 0.0,
        "total_bottlenecks": total_bottlenecks,
        "delayed_orders_count": delayed_count,
        "adjusted_operations": adjusted_count,
        "mrp_ticket": ticket.mrp_ticket or "",
        "execution_time": ticket.execution_time_seconds or 0.0
    }

@frappe.whitelist()
def get_scheduled_operations_summary(ticket_name, workstation=None, work_order=None, status=None, start=0, page_length=50):
    """Returns paginated scheduled operations for the detailed grid."""
    filters = {"aps_ticket": ticket_name}
    if workstation:
        filters["workstation"] = workstation
    if work_order:
        filters["work_order"] = ["like", f"%{work_order}%"]
    if status:
        filters["status"] = status

    total_count = frappe.db.count("APS Scheduled Operation", filters=filters)
    items = frappe.get_all(
        "APS Scheduled Operation",
        filters=filters,
        fields=[
            "name", "work_order", "mrp_ticket", "production_item", "item_name",
            "operation", "sequence_id", "workstation", "planned_start_time",
            "planned_end_time", "duration_mins", "setup_mins", "runtime_mins",
            "qty_to_produce", "is_adjusted", "has_conflict", "status", "delay_hours"
        ],
        start=cint(start),
        page_length=cint(page_length),
        order_by="planned_start_time ASC, sequence_id ASC"
    )

    return {
        "total": total_count,
        "items": items
    }

def style_excel_sheet(ws, title, headers, rows, header_color="1B365D"):
    """Applies professional enterprise styling to an openpyxl worksheet."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    # Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = title
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="0F2027", end_color="0F2027", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28

    ws.append([]) # row 2 spacer
    ws.row_dimensions[2].height = 6

    ws.append(headers)
    ws.row_dimensions[3].height = 24

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=3, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

    current_row = 4
    for r in rows:
        ws.append(r)
        fill = zebra_fill if (current_row % 2 == 0) else white_fill
        ws.row_dimensions[current_row].height = 20

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

            val = cell.value
            if isinstance(val, (int, float)):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

@frappe.whitelist()
def export_aps_gantt_excel(ticket_name):
    """
    Exports the updated Gantt scheduled operations in chronological sequence and formatted in list mode.
    """
    import openpyxl
    import io
    from frappe.desk.utils import provide_binary_file

    wb = openpyxl.Workbook()

    # Sheet 1: Sequenciamento Gantt
    ws1 = wb.active
    ws1.title = "Sequenciamento Gantt"

    ops = frappe.get_all(
        "APS Scheduled Operation",
        filters={"aps_ticket": ticket_name},
        fields=[
            "work_order", "mrp_ticket", "production_item", "item_name",
            "operation", "sequence_id", "workstation", "planned_start_time",
            "planned_end_time", "duration_mins", "setup_mins", "runtime_mins",
            "qty_to_produce", "predecessor_operation", "is_adjusted", "status", "delay_hours"
        ],
        order_by="planned_start_time ASC, sequence_id ASC"
    )

    title1 = f"ERPZ APS — Programação e Sequenciamento da Produção (Gantt) | Ticket: {ticket_name}"
    headers1 = [
        "Ordem Cronológica", "Início Programado", "Término Programado", "Posto de Trabalho",
        "Ordem de Produção", "Código Item", "Descrição do Produto", "Operação", "Seq Roteiro",
        "Duração (min)", "Setup (min)", "Processamento (min)", "Quantidade",
        "Operação Predecessora", "Situação", "Ajustado no Gantt", "Atraso (h)", "Ticket MRP Origem"
    ]
    rows1 = []
    for idx, o in enumerate(ops, start=1):
        rows1.append([
            f"#{idx}",
            str(o.planned_start_time) if o.planned_start_time else "",
            str(o.planned_end_time) if o.planned_end_time else "",
            o.workstation,
            o.work_order,
            o.production_item,
            o.item_name or "",
            o.operation,
            o.sequence_id,
            flt(o.duration_mins),
            flt(o.setup_mins),
            flt(o.runtime_mins),
            flt(o.qty_to_produce),
            o.predecessor_operation or "Nenhuma (Primeira)",
            o.status,
            "SIM (Ajuste Manual)" if o.is_adjusted else "Não",
            flt(o.delay_hours),
            o.mrp_ticket or ""
        ])

    style_excel_sheet(ws1, title1, headers1, rows1, header_color="1B365D")

    # Sheet 2: Carga x Capacidade (CRP)
    ws2 = wb.create_sheet(title="Carga x Capacidade (CRP)")
    crp_rows = frappe.get_all(
        "APS Capacity Load",
        filters={"aps_ticket": ticket_name},
        fields=[
            "workstation", "workstation_name", "period_date",
            "available_capacity_hours", "planned_load_hours",
            "capacity_balance_hours", "utilization_pct", "is_bottleneck",
            "overload_hours", "operations_count"
        ],
        order_by="period_date ASC, workstation ASC"
    )
    title2 = f"ERPZ APS — Carga x Capacidade por Posto (CRP) | Ticket: {ticket_name}"
    headers2 = [
        "Posto de Trabalho", "Descrição", "Data", "Capacidade Disponível (h)",
        "Carga Programada (h)", "Saldo Capacidade (h)", "Utilização (%)",
        "Sobrecarga (h)", "Operações Alocadas", "Situação"
    ]
    rows2 = []
    for c in crp_rows:
        rows2.append([
            c.workstation,
            c.workstation_name or "",
            str(c.period_date),
            flt(c.available_capacity_hours),
            flt(c.planned_load_hours),
            flt(c.capacity_balance_hours),
            flt(c.utilization_pct),
            flt(c.overload_hours),
            c.operations_count,
            "GARGALO / SOBRECARGA" if c.is_bottleneck else "Normal"
        ])
    style_excel_sheet(ws2, title2, headers2, rows2, header_color="2C5282")

    buf = io.BytesIO()
    wb.save(buf)
    provide_binary_file(f"Gantt_Sequenciamento_{ticket_name}", "xlsx", buf.getvalue())

@frappe.whitelist()
def download_routing_template(item_code=None):
    """Generates an Excel template for bulk importing product operations routing."""
    import openpyxl
    import io
    from frappe.desk.utils import provide_binary_file

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roteiro Produtivo"

    title = "ERPZ APS — Modelo de Importação do Roteiro Produtivo por Produto"
    headers = [
        "Código do Item*", "Sequência*", "Código da Operação*", "Nome da Operação",
        "Estação de Trabalho*", "Qtd Referência*", "Tempo Referência (min)*",
        "Setup (min)", "Transferência (min)", "Sobreposição Próx Op (%)",
        "Sobreposição Entre OPs (%)", "Estação Alternativa Produto", "Instruções Técnicas"
    ]
    code = item_code or "TEST-MRP-A"
    rows = [
        [code, 10, "10 - Cortar", "Cortar", "SERRA 001", 5.0, 15.0, 10.0, 15.0, 50.0, 20.0, "Centro Usinagem 02", "Corte longitudinal"],
        [code, 20, "20 - Dobrar", "Dobrar", "Centro Usinagem 02", 10.0, 20.0, 15.0, 15.0, 30.0, 10.0, "Torno CNC 01", "Dobra em 90 graus"],
        [code, 30, "30 - Embalar", "Embalar", "Torno CNC 01", 50.0, 25.0, 5.0, 0.0, 0.0, 0.0, "", "Embalagem e identificação"]
    ]

    style_excel_sheet(ws, title, headers, rows, header_color="1B365D")

    buf = io.BytesIO()
    wb.save(buf)
    provide_binary_file("Modelo_Roteiro_Processo_Produtivo_APS", "xlsx", buf.getvalue())

@frappe.whitelist()
def import_product_routing_excel(file_url=None):
    """
    Imports product operations routing from an uploaded Excel file,
    populating or updating APS Product Routing documents.
    """
    import openpyxl
    import io
    from collections import defaultdict

    if not frappe.has_permission("APS Product Routing", "write"):
        frappe.throw(_("Sem permissão para cadastrar Roteiros Produtivos."))

    file_content = None
    if "file" in frappe.request.files:
        file_content = frappe.request.files["file"].read()
    elif file_url:
        _file = frappe.get_doc("File", {"file_url": file_url})
        file_content = _file.get_content()

    if not file_content:
        frappe.throw(_("Nenhum arquivo enviado para importação."))

    wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    ws = wb.active

    # Find header row
    header_row_idx = 1
    for r_idx in range(1, 10):
        val = str(ws.cell(row=r_idx, column=1).value or "").lower()
        if "código" in val or "codigo" in val or "item" in val:
            header_row_idx = r_idx
            break

    start_data_row = header_row_idx + 1

    # Group rows by item_code
    rows_by_item = defaultdict(list)
    errors = []

    for r_num in range(start_data_row, ws.max_row + 1):
        item_code = str(ws.cell(row=r_num, column=1).value or "").strip()
        if not item_code:
            continue

        raw_seq = ws.cell(row=r_num, column=2).value
        op_code = str(ws.cell(row=r_num, column=3).value or "").strip()
        op_name = str(ws.cell(row=r_num, column=4).value or "").strip()
        workstation = str(ws.cell(row=r_num, column=5).value or "").strip()
        raw_batch = ws.cell(row=r_num, column=6).value
        raw_cycle = ws.cell(row=r_num, column=7).value
        raw_setup = ws.cell(row=r_num, column=8).value
        raw_trans = ws.cell(row=r_num, column=9).value
        raw_overlap = ws.cell(row=r_num, column=10).value
        raw_wo_overlap = ws.cell(row=r_num, column=11).value
        alt_ws = str(ws.cell(row=r_num, column=12).value or "").strip()
        desc = str(ws.cell(row=r_num, column=13).value or "").strip()

        if not frappe.db.exists("Item", item_code):
            errors.append(f"Linha {r_num}: Item '{item_code}' não encontrado.")
            continue

        if not frappe.db.exists("Workstation", workstation):
            errors.append(f"Linha {r_num}: Estação de trabalho '{workstation}' não encontrada.")
            continue

        batch_qty = flt(raw_batch) if raw_batch else 1.0
        cycle_time = flt(raw_cycle) if raw_cycle else 10.0
        unit_time = round(cycle_time / max(0.001, batch_qty), 4)

        rows_by_item[item_code].append({
            "sequence_id": cint(raw_seq or 10),
            "operation_code": op_code or f"Op {raw_seq}",
            "operation_name": op_name or op_code,
            "workstation": workstation,
            "batch_qty": batch_qty,
            "cycle_time_mins": cycle_time,
            "time_per_unit_mins": unit_time,
            "setup_time_mins": flt(raw_setup or 0.0),
            "transfer_time_mins": flt(raw_trans or 15.0),
            "overlap_pct": flt(raw_overlap or 0.0),
            "allow_wo_overlap_pct": flt(raw_wo_overlap or 0.0),
            "alternative_workstation": alt_ws if frappe.db.exists("Workstation", alt_ws) else None,
            "description": desc
        })

    imported_routings = 0
    total_ops = 0

    for item_code, op_list in rows_by_item.items():
        if frappe.db.exists("APS Product Routing", {"item_code": item_code}):
            routing = frappe.get_doc("APS Product Routing", {"item_code": item_code})
        else:
            routing = frappe.new_doc("APS Product Routing")
            routing.item_code = item_code

        routing.routing_name = f"Roteiro Padrão - {item_code}"
        routing.is_active = 1
        routing.operations = []

        # Sort by sequence
        op_list.sort(key=lambda x: x["sequence_id"])

        for op_data in op_list:
            routing.append("operations", op_data)
            total_ops += 1

        routing.save(ignore_permissions=True)
        imported_routings += 1

    frappe.db.commit()

    return {
        "success": True,
        "imported_routings": imported_routings,
        "total_operations": total_ops,
        "errors": errors
    }
